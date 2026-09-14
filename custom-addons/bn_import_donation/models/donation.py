from odoo import models, fields, api, _
from odoo.exceptions import ValidationError


state_selection = [
    ('draft', 'Draft'),
    ('posted', 'Posted')
]


class Donation(models.Model):
    _name = 'donation'
    _description = "Donation"
    _inherit = ["mail.thread", "mail.activity.mixin"]


    donor_id = fields.Many2one('res.partner', string="Donor / Student", tracking=True)
    journal_id = fields.Many2one('account.journal', string="Journal", tracking=True)
    product_id = fields.Many2one('product.product', string="Product", tracking=True)
    gateway_config_id = fields.Many2one('gateway.config', string="Gateway Config", tracking=True)
    company_id = fields.Many2one('res.company', string="Company", default=lambda self: self.env.user.company_id.id)
    currency_id = fields.Many2one(related='company_id.currency_id', string="Currency")
    import_donation_id = fields.Many2one('import.donation', string="Import Donation")

    name = fields.Char('Name', default="New", tracking=True)
    transaction_id = fields.Char('Transaction ID', tracking=True)

    reference = fields.Text('Reference/Remarks', tracking=True)
    
    date = fields.Char('Date', tracking=True)

    amount = fields.Monetary('Amount', tracking=True)

    is_fee = fields.Boolean('Is Fee', tracking=True)

    state = fields.Selection(selection=state_selection, string="State", default="draft", tracking=True)


    @api.model
    def create(self, vals):
        if vals.get('name', _('New') == _('New')):
            vals['name'] = self.env['ir.sequence'].next_by_code('import_donation') or ('New')

        return super(Donation, self).create(vals)
    
    def action_confirm(self):
        self.state = 'posted'
    
    def action_draft(self):
        self.state = 'draft'

    def action_sync_existing_donors(self):
        import base64
        from io import BytesIO
        import openpyxl
        import xlrd

        if not self:
            raise ValidationError(_("No donation records selected. Please select one or more donations first."))

        Partner = self.env['res.partner']
        GLITCH_NAME = '3 START KK MART'

        fixed_count = 0
        skipped_count = 0

        # cache parsed excel name_map per import_donation_id, so we don't re-parse per row
        name_map_cache = {}

        def get_name_map(import_donation):
            if import_donation.id in name_map_cache:
                return name_map_cache[import_donation.id]

            name_map = {}
            if import_donation.import_file:
                data = base64.b64decode(import_donation.import_file)
                stream = BytesIO(data)

                if data[:4] == b'PK\x03\x04':
                    workbook = openpyxl.load_workbook(stream)
                    sheet = workbook.active
                    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
                    rows = sheet.iter_rows(min_row=2, values_only=True)
                elif data[:4] == b'\xD0\xCF\x11\xE0':
                    workbook = xlrd.open_workbook(file_contents=data)
                    sheet = workbook.sheet_by_index(0)
                    headers = sheet.row_values(0)
                    rows = (sheet.row_values(i) for i in range(1, sheet.nrows))
                else:
                    rows = []
                    headers = []

                if 'Id' in headers and 'From name' in headers:
                    id_idx = headers.index('Id')
                    name_idx = headers.index('From name')
                    for row in rows:
                        raw_id = row[id_idx]
                        if isinstance(raw_id, float) and raw_id.is_integer():
                            txn_id = str(int(raw_id))
                        else:
                            txn_id = str(raw_id or '').strip()
                        from_name = row[name_idx]
                        if txn_id and from_name:
                            name_map[txn_id] = from_name

            name_map_cache[import_donation.id] = name_map
            return name_map

        for donation in self:
            if not donation.donor_id or (donation.donor_id.name or '').strip() != GLITCH_NAME:
                skipped_count += 1
                continue
            txn_id = (donation.transaction_id or '').strip()

            # 1st try: donor_student_name already on the valid line
            source_line = self.env['valid.import.donation'].search([
                ('transaction_id', '=', txn_id)
            ], limit=1)

            correct_name = False
            if source_line and source_line.donor_student_name:
                correct_name = source_line.donor_student_name.strip()

            # 2nd try: fall back to reading the Excel file directly
            if not correct_name:
                import_donation = donation.import_donation_id or (source_line.import_donation_id if source_line else False)
                if import_donation:
                    name_map = get_name_map(import_donation)
                    correct_name = name_map.get(txn_id)

            if not correct_name or correct_name == GLITCH_NAME:
                skipped_count += 1
                continue

            partner = Partner.search([('name', '=', correct_name)], limit=1)

            if not partner:
                partner = Partner.create({
                    'name': correct_name,
                    'mobile': source_line.mobile if source_line else False,
                    'cnic_no': source_line.cnic_no if source_line else False,
                    'email': source_line.email if source_line else False,
                })

            donation.donor_id = partner.id

            # also backfill the valid.import.donation line, so future runs don't hit this again
            if source_line and not source_line.donor_student_name:
                source_line.donor_student_name = correct_name

            fixed_count += 1

        if fixed_count == 0:
            raise ValidationError(
                _("No donations were fixed. None of the selected records currently "
                  "show '%s', or no name could be found in the linked import line "
                  "or the uploaded Excel file. %s record(s) were skipped.")
                % (GLITCH_NAME, skipped_count)
            )

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Donor Name Fixed'),
                'message': _('%s donation(s) fixed, %s skipped.') % (fixed_count, skipped_count),
                'type': 'success',
                'sticky': False,
            },
        }