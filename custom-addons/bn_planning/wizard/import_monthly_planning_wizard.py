from odoo import models, fields, api
from odoo.exceptions import ValidationError
from datetime import datetime, date
import base64
from io import BytesIO
import logging

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    openpyxl = None


TAB_MODEL_MAP = {
    'kitchen':   'monthly.planning.kitchen',
    'madaris':   'monthly.planning.madaris',
    'medical':   'monthly.planning.medical',
    'livestock': 'monthly.planning.livestock',
    'food':      'monthly.planning.food',
    'ration':    'monthly.planning.ration',
    'meat':      'monthly.planning.meat',
}


class ImportMonthlyPlanningWizard(models.TransientModel):
    _name = 'import.monthly.planning.wizard'
    _description = 'Import Monthly Planning from Excel'

    monthly_planning_id = fields.Many2one(
        'monthly.planning', string='Monthly Planning',
        required=True, readonly=True,
    )
    file_data = fields.Binary(string='Excel File', required=True)
    file_name = fields.Char(string='File Name')
    clear_existing = fields.Boolean(
        string='Delete existing lines before import',
        default=True,
        help='If enabled, all existing lines of the enabled tabs will be removed before import.',
    )

    # ─────────────────── helpers ───────────────────
    def _parse_date(self, value):
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(str(value).strip(), fmt).date()
            except (ValueError, TypeError):
                continue
        return False

    def _find_product(self, value):
        if not value:
            return False
        Product = self.env['product.product']
        val = str(value).strip()
        product = Product.search([
            '|', '|',
            ('default_code', '=', val),
            ('barcode', '=', val),
            ('name', '=', val),
        ], limit=1)
        if not product:
            product = Product.search([('name', 'ilike', val)], limit=1)
        return product

    def _match_sheet_to_tab(self, sheet_name):
        sname = (sheet_name or '').strip().lower()
        sname = sname.replace(' planning', '').replace('_planning', '').strip()
        return sname if sname in TAB_MODEL_MAP else None

    # ─────────────────── main action ───────────────────
    def action_import(self):
        self.ensure_one()

        if not openpyxl:
            raise ValidationError(
                "The 'openpyxl' Python library is required to import Excel files.\n"
                "Install it with: pip install openpyxl"
            )

        mp = self.monthly_planning_id

        # Read the file
        try:
            raw = base64.b64decode(self.file_data)
            wb = openpyxl.load_workbook(BytesIO(raw), data_only=True)
        except Exception as e:
            raise ValidationError(f"Could not read the Excel file:\n{e}")

        # Which tabs are enabled on the Planning Type?
        enabled_tabs = [
            tab for tab in TAB_MODEL_MAP
            if mp.planning_type_id and getattr(mp.planning_type_id, tab, False)
        ]
        if not enabled_tabs:
            raise ValidationError(
                "No tabs are enabled on the selected Planning Type."
            )

        # Optionally clear existing lines
        if self.clear_existing:
            for tab in enabled_tabs:
                getattr(mp, f"{tab}_line_ids").unlink()

        created = updated = skipped = 0
        errors = []

        for sheet_name in wb.sheetnames:
            tab_key = self._match_sheet_to_tab(sheet_name)
            if not tab_key:
                _logger.info("Skipping sheet '%s' (no matching tab).", sheet_name)
                continue
            if tab_key not in enabled_tabs:
                _logger.info("Skipping sheet '%s' (tab not enabled).", sheet_name)
                continue

            Model = self.env[TAB_MODEL_MAP[tab_key]]
            sheet = wb[sheet_name]

            # Row 1 is the header — start at row 2
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not row or all(c is None or str(c).strip() == '' for c in row):
                    continue

                date_val    = row[0] if len(row) > 0 else None
                product_val = row[1] if len(row) > 1 else None
                qty_val     = row[2] if len(row) > 2 else None
                location_val = row[3] if len(row) > 3 else None

                if not product_val:
                    skipped += 1
                    continue

                date_obj = self._parse_date(date_val)
                if not date_obj:
                    errors.append(f"[{sheet_name}] Row {row_idx}: invalid date '{date_val}'")
                    skipped += 1
                    continue

                product = self._find_product(product_val)
                if not product:
                    errors.append(f"[{sheet_name}] Row {row_idx}: product not found '{product_val}'")
                    skipped += 1
                    continue

                try:
                    quantity = float(qty_val) if qty_val not in (None, '') else 0.0
                except (ValueError, TypeError):
                    quantity = 0.0

                vals = {
                    'monthly_planning_id': mp.id,
                    'date': date_obj,
                    'product_id': product.id,
                    'quantity': quantity,
                }

                # Optional location (for livestock)
                if tab_key == 'livestock' and location_val:
                    loc = self.env['stock.location'].search(
                        [('name', 'ilike', str(location_val).strip())], limit=1
                    )
                    if loc:
                        vals['location_id'] = loc.id

                existing = Model.search([
                    ('monthly_planning_id', '=', mp.id),
                    ('date', '=', date_obj),
                    ('product_id', '=', product.id),
                ], limit=1)

                if existing:
                    existing.write({'quantity': quantity})
                    updated += 1
                else:
                    Model.create(vals)
                    created += 1

        # Build result message
        lines = [
            f"Created: {created}",
            f"Updated: {updated}",
            f"Skipped: {skipped}",
        ]
        if errors:
            lines.append("")
            lines.append("Errors:")
            lines.extend(errors[:15])
            if len(errors) > 15:
                lines.append(f"... and {len(errors) - 15} more")

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Excel Import Complete',
                'message': '\n'.join(lines),
                'type': 'warning' if errors else 'success',
                'sticky': bool(errors),
                'next': {'type': 'ir.actions.client', 'tag': 'reload'},
            }
        }